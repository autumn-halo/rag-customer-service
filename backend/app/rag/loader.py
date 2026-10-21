"""文档解析：把上传文件统一解析成 [ParsedPage] 中间结构。

- PDF：逐页提取文字；提取为空 → 判定疑似扫描版，报可读错误（OCR 列入后续规划）
- Word：段落 + 表格都提取（表格行拼为 "列名 | 值"）
- Excel：每行生成 "列名: 值；列名: 值" 文本，sheet 名进 metadata
- TXT：自动探测 GBK/UTF-8 等编码
- MD：去除语法符号，保留标题层级文本
"""
import re
from dataclasses import dataclass, field
from pathlib import Path

from charset_normalizer import from_path
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_TYPES = {"pdf", "docx", "xlsx", "txt", "md"}


@dataclass
class ParsedPage:
    page_no: int
    text: str
    metadata: dict = field(default_factory=dict)


class DocumentParseError(Exception):
    """文档无法解析（损坏文件、扫描版 PDF、空文档等），信息可直接展示给用户。"""


def parse_document(path: Path, file_type: str) -> list[ParsedPage]:
    if file_type == "pdf":
        return _parse_pdf(path)
    if file_type == "docx":
        return _parse_docx(path)
    if file_type == "xlsx":
        return _parse_xlsx(path)
    if file_type == "txt":
        return _parse_txt(path)
    if file_type == "md":
        return _parse_md(path)
    raise DocumentParseError(f"不支持的文件类型: {file_type}")


def _parse_pdf(path: Path) -> list[ParsedPage]:
    try:
        reader = PdfReader(str(path))
    except Exception as e:  # noqa: BLE001
        raise DocumentParseError(f"PDF 文件损坏或无法打开: {e}") from e
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(ParsedPage(page_no=i, text=text))
    if not pages:
        raise DocumentParseError("未能从 PDF 中提取到文字：疑似扫描版 PDF，暂不支持 OCR 识别")
    return pages


def _parse_docx(path: Path) -> list[ParsedPage]:
    try:
        doc = DocxDocument(str(path))
    except Exception as e:  # noqa: BLE001
        raise DocumentParseError(f"Word 文件损坏或无法打开: {e}") from e
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text:
        raise DocumentParseError("Word 文档中没有可提取的文字内容")
    return [ParsedPage(page_no=1, text=text)]


def _parse_xlsx(path: Path) -> list[ParsedPage]:
    try:
        wb = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise DocumentParseError(f"Excel 文件损坏或无法打开: {e}") from e
    pages = []
    for ws in wb.worksheets:
        lines = []
        header: list[str] | None = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if not any(cells):
                continue
            if header is None:
                header = cells
                continue
            pairs = [f"{h}: {v}" for h, v in zip(header, cells) if v]
            if pairs:
                lines.append("；".join(pairs))
        if lines:
            pages.append(
                ParsedPage(
                    page_no=len(pages) + 1,
                    text="\n".join(lines),
                    metadata={"sheet": ws.title},
                )
            )
    if not pages:
        raise DocumentParseError("Excel 表格中没有数据")
    return pages


def _parse_txt(path: Path) -> list[ParsedPage]:
    best = from_path(str(path)).best()
    if best is None:
        raise DocumentParseError("无法识别文本文件编码")
    text = str(best).strip()
    if not text:
        raise DocumentParseError("文本文件为空")
    return [ParsedPage(page_no=1, text=text)]


def _parse_md(path: Path) -> list[ParsedPage]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    text = re.sub(r"^#{1,6}\s+", "", raw, flags=re.MULTILINE)  # 标题符号
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)  # 加粗
    text = re.sub(r"`(.+?)`", r"\1", text)  # 行内代码
    text = text.strip()
    if not text:
        raise DocumentParseError("Markdown 文件为空")
    return [ParsedPage(page_no=1, text=text)]
