"""文档解析测试：动态生成五种格式的样例文件（不提交二进制到 git）。"""
import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from app.rag.loader import DocumentParseError, parse_document


def test_pdf_parsing(tmp_path):
    path = tmp_path / "test.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    c = canvas.Canvas(str(path))
    c.setFont("STSong-Light", 14)
    c.drawString(50, 800, "Return policy: refund within 7 days.")
    c.save()

    pages = parse_document(path, "pdf")
    assert len(pages) == 1
    assert "refund within 7 days" in pages[0].text
    assert pages[0].page_no == 1


def test_corrupted_pdf_raises(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"this is not a real pdf file")
    with pytest.raises(DocumentParseError, match="PDF"):
        parse_document(path, "pdf")


def test_docx_paragraphs_and_tables(tmp_path):
    path = tmp_path / "test.docx"
    doc = DocxDocument()
    doc.add_paragraph("商品参数说明")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "型号"
    table.rows[0].cells[1].text = "A100"
    table.rows[1].cells[0].text = "颜色"
    table.rows[1].cells[1].text = "黑色"
    doc.save(str(path))

    pages = parse_document(path, "docx")
    assert "商品参数说明" in pages[0].text
    assert "型号 | A100" in pages[0].text
    assert "颜色 | 黑色" in pages[0].text


def test_xlsx_rows_become_key_value_lines(tmp_path):
    path = tmp_path / "test.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "商品表"
    ws.append(["品名", "价格"])
    ws.append(["蓝牙耳机", "199"])
    wb.save(str(path))

    pages = parse_document(path, "xlsx")
    assert pages[0].metadata["sheet"] == "商品表"
    assert "品名: 蓝牙耳机；价格: 199" in pages[0].text


def test_gbk_txt_decoded(tmp_path):
    path = tmp_path / "gbk.txt"
    path.write_bytes("退货政策说明：七天无理由退货。".encode("gbk"))

    pages = parse_document(path, "txt")
    assert "七天无理由退货" in pages[0].text


def test_empty_txt_raises(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_bytes(b"")
    with pytest.raises(DocumentParseError):
        parse_document(path, "txt")


def test_markdown_syntax_stripped(tmp_path):
    path = tmp_path / "faq.md"
    path.write_text("# 常见问题\n**发货时间**一般为 48 小时。\n", encoding="utf-8")

    pages = parse_document(path, "md")
    text = pages[0].text
    assert "常见问题" in text
    assert "发货时间" in text
    assert "#" not in text
    assert "**" not in text


def test_unsupported_type_raises(tmp_path):
    with pytest.raises(DocumentParseError, match="不支持"):
        parse_document(tmp_path / "x.exe", "exe")
